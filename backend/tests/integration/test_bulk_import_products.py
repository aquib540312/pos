import io

import openpyxl

from app.modules.catalog.bulk_import import ALL_COLUMNS


def _workbook_bytes(rows: list[dict]) -> bytes:
    """Builds each row from a {column_name: value} dict rather than a
    positional list, so the test data can't silently drift out of sync
    with ALL_COLUMNS' actual column order."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(ALL_COLUMNS)
    for row in rows:
        sheet.append([row.get(col) for col in ALL_COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(client, headers, content: bytes, filename: str = "products.xlsx"):
    return client.post(
        "/api/v1/catalog/products/bulk-import",
        headers=headers,
        files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_bulk_import_creates_new_products(client, seeded_org):
    content = _workbook_bytes(
        [
            {"sku": "MLK-001", "name": "Toned Milk 500ml", "uom_code": "PCS", "mrp": 30, "sale_price": 28, "barcode": "8900000000001", "tracks_expiry": "TRUE"},
            {"sku": "EGG-006", "name": "Eggs (6 pack)", "uom_code": "PCS", "mrp": 60, "sale_price": 55},
        ]
    )
    resp = _upload(client, seeded_org["auth_headers"], content)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["total"] == 2
    assert body["created"] == 2
    assert body["updated"] == 0
    assert body["failed"] == 0

    products = client.get("/api/v1/catalog/products", headers=seeded_org["auth_headers"]).json()
    skus = {p["sku"] for p in products}
    assert {"MLK-001", "EGG-006"} <= skus


def test_bulk_import_updates_existing_product_by_sku(client, seeded_org):
    """seeded_org already has a BRD-001 product (name "White Bread 400g",
    sale_price 40) -- a row with that SKU should update it in place, not
    create a duplicate."""
    content = _workbook_bytes(
        [
            {
                "sku": "BRD-001",
                "name": "White Bread 400g (Family Pack)",
                "uom_code": "PCS",
                "mrp": 55,
                "sale_price": 50,
                "barcode": "8901234567890",
                "hsn_code": "1905",
                "tracks_batches": "TRUE",
                "tracks_expiry": "TRUE",
            }
        ]
    )
    resp = _upload(client, seeded_org["auth_headers"], content)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["created"] == 0
    assert body["updated"] == 1
    assert body["rows"][0]["status"] == "updated"

    products = client.get("/api/v1/catalog/products", headers=seeded_org["auth_headers"]).json()
    updated = next(p for p in products if p["sku"] == "BRD-001")
    assert updated["name"] == "White Bread 400g (Family Pack)"
    assert updated["sale_price"] == 50


def test_bulk_import_reports_row_error_without_discarding_valid_rows(client, seeded_org):
    content = _workbook_bytes(
        [
            {"sku": "GOOD-001", "name": "A Valid Product", "uom_code": "PCS", "mrp": 20, "sale_price": 18},
            {"sku": "BAD-001", "name": "Unknown Unit Product", "uom_code": "NOPE", "mrp": 20, "sale_price": 18},
        ]
    )
    resp = _upload(client, seeded_org["auth_headers"], content)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["total"] == 2
    assert body["created"] == 1
    assert body["failed"] == 1
    error_row = next(r for r in body["rows"] if r["status"] == "error")
    assert error_row["sku"] == "BAD-001"
    assert "NOPE" in error_row["error"]

    products = client.get("/api/v1/catalog/products", headers=seeded_org["auth_headers"]).json()
    skus = {p["sku"] for p in products}
    assert "GOOD-001" in skus
    assert "BAD-001" not in skus


def test_bulk_import_auto_creates_category_by_name(client, seeded_org):
    content = _workbook_bytes(
        [{"sku": "SNK-001", "name": "Potato Chips", "uom_code": "PCS", "mrp": 20, "sale_price": 18, "category": "Snacks"}]
    )
    resp = _upload(client, seeded_org["auth_headers"], content)
    assert resp.status_code == 200, resp.text
    assert resp.json()["created"] == 1

    categories = client.get("/api/v1/catalog/categories", headers=seeded_org["auth_headers"]).json()
    assert any(c["name"] == "Snacks" for c in categories)


def test_bulk_import_rejects_workbook_missing_required_columns(client, seeded_org):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["sku", "name"])  # missing uom_code, mrp, sale_price
    sheet.append(["X-1", "Something"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    resp = _upload(client, seeded_org["auth_headers"], buffer.getvalue())
    assert resp.status_code == 422
    assert "uom_code" in resp.json()["detail"]


def test_bulk_import_rejects_non_excel_file(client, seeded_org):
    resp = _upload(client, seeded_org["auth_headers"], b"not an excel file", filename="products.txt")
    assert resp.status_code == 422


def test_download_bulk_import_template_has_expected_headers(client, seeded_org):
    resp = client.get("/api/v1/catalog/products/bulk-import/template", headers=seeded_org["auth_headers"])
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    workbook = openpyxl.load_workbook(io.BytesIO(resp.content))
    sheet = workbook["Products"]
    header = [c.value for c in sheet[1]]
    assert header == ALL_COLUMNS

    # The example row's values must line up with their own column, not just
    # exist in *some* column -- catches the header/value order drifting
    # apart from ALL_COLUMNS (this happened for real: mrp ended up holding
    # a category name).
    example_row = dict(zip(header, (c.value for c in sheet[2]), strict=True))
    assert example_row["uom_code"] == "PCS"
    assert isinstance(example_row["mrp"], (int, float))
    assert isinstance(example_row["sale_price"], (int, float))
