"""Parsing and coercion for the bulk product upload (Excel) feature. Kept
separate from service.py because it's file-format concerns (openpyxl,
header matching, cell coercion), not catalog business logic -- the DB
lookups/upsert-by-SKU orchestration that actually uses this lives in
CatalogService.bulk_import_products.
"""

import io
from dataclasses import dataclass, field

import openpyxl

from app.core.exceptions import ValidationError

REQUIRED_COLUMNS = ["sku", "name", "uom_code", "mrp", "sale_price"]
OPTIONAL_COLUMNS = [
    "barcode",
    "description",
    "category",
    "hsn_code",
    "purchase_price",
    "reorder_level",
    "tracks_batches",
    "tracks_serials",
    "tracks_expiry",
]
ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS

# Keeps one upload request bounded regardless of how large a file someone
# throws at it -- same reasoning as sync_push_max_batch_size elsewhere.
MAX_ROWS = 2000

_TRUE_VALUES = {"true", "1", "yes", "y"}


def parse_product_workbook(file_bytes: bytes) -> list[tuple[int, dict[str, object]]]:
    """Row 1 is the header (case/whitespace-insensitive); every following
    non-blank row is one product. Returns (spreadsheet_row_number, {column:
    value}) pairs so a validation error can point the uploader back at the
    exact row to fix."""
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError(f"Could not read the uploaded file as an Excel (.xlsx) workbook: {exc}") from exc

    sheet = workbook.worksheets[0]
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ValidationError("The uploaded workbook has no rows") from None

    header = [str(cell).strip().lower() if cell is not None else "" for cell in header_row]
    missing = [col for col in REQUIRED_COLUMNS if col not in header]
    if missing:
        raise ValidationError(f"Missing required column(s): {', '.join(missing)}")

    rows: list[tuple[int, dict[str, object]]] = []
    for row_number, raw_row in enumerate(rows_iter, start=2):
        if all(cell is None for cell in raw_row):
            continue  # blank spacer row -- common at the end of a sheet
        if len(rows) >= MAX_ROWS:
            raise ValidationError(f"Too many rows -- max {MAX_ROWS} products per upload")
        rows.append((row_number, dict(zip(header, raw_row, strict=False))))
    return rows


def generate_product_import_template() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(ALL_COLUMNS)
    example_row = {
        "sku": "BRD-002",
        "name": "Brown Bread 400g",
        "uom_code": "PCS",
        "mrp": 50,
        "sale_price": 45,
        "barcode": "8901234567891",
        "description": None,
        "category": "Bakery",
        "hsn_code": "1905",
        "purchase_price": 32,
        "reorder_level": 10,
        "tracks_batches": "FALSE",
        "tracks_serials": "FALSE",
        "tracks_expiry": "TRUE",
    }
    sheet.append([example_row[col] for col in ALL_COLUMNS])

    notes = workbook.create_sheet("Instructions")
    notes.append(["Column", "Required", "Notes"])
    notes.append(["sku", "Yes", "Unique per organization. A SKU that already exists updates that product instead of creating a new one."])
    notes.append(["name", "Yes", "Product display name."])
    notes.append(["uom_code", "Yes", "Must already exist under Catalog > Units of Measure (e.g. PCS, KG)."])
    notes.append(["mrp", "Yes", "Maximum retail price."])
    notes.append(["sale_price", "Yes", "Selling price."])
    notes.append(["barcode", "No", "Leave blank if the product has no barcode."])
    notes.append(["category", "No", "Auto-created if it doesn't already exist."])
    notes.append(["hsn_code", "No", "Must already exist under Catalog > HSN/SAC Codes if set."])
    notes.append(["purchase_price", "No", "Defaults to 0."])
    notes.append(["reorder_level", "No", "Defaults to 0."])
    notes.append(["tracks_batches / tracks_serials / tracks_expiry", "No", "TRUE or FALSE. Defaults to FALSE."])
    notes.append(["", "", "Combo products aren't supported by bulk import -- create those individually."])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def require_str(row: dict[str, object], key: str) -> str:
    value = row.get(key)
    if value is None or str(value).strip() == "":
        raise ValidationError(f"'{key}' is required")
    return str(value).strip()


def optional_str(row: dict[str, object], key: str) -> str | None:
    value = row.get(key)
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def require_number(row: dict[str, object], key: str) -> float:
    value = row.get(key)
    if value is None or str(value).strip() == "":
        raise ValidationError(f"'{key}' is required")
    try:
        return float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        raise ValidationError(f"'{key}' must be a number, got '{value}'") from None


def optional_number(row: dict[str, object], key: str, default: float) -> float:
    value = row.get(key)
    if value is None or str(value).strip() == "":
        return default
    try:
        return float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        raise ValidationError(f"'{key}' must be a number, got '{value}'") from None


def optional_bool(row: dict[str, object], key: str, default: bool) -> bool:
    value = row.get(key)
    if value is None or str(value).strip() == "":
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in _TRUE_VALUES


@dataclass
class RowResult:
    row_number: int
    status: str  # "created" | "updated" | "error"
    sku: str | None
    error: str | None = None


@dataclass
class BulkImportResult:
    results: list[RowResult] = field(default_factory=list)

    @property
    def created(self) -> int:
        return sum(1 for r in self.results if r.status == "created")

    @property
    def updated(self) -> int:
        return sum(1 for r in self.results if r.status == "updated")

    @property
    def failed(self) -> int:
        return sum(1 for r in self.results if r.status == "error")
